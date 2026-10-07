#!/usr/bin/env python3
"""
Áp định dạng của form mẫu (Nhóm 5) lên file DOCX do pandoc sinh ra.

Form mẫu quy định: Times New Roman, thân bài 13pt giãn dòng 1,5, thụt đầu dòng 1cm,
căn đều hai biên; chú thích hình và bảng 12pt in đậm; ô bảng 11pt; trang bìa căn giữa.
Script cũng thay mục lục viết tay bằng trường TOC của Word để Word tự đánh số trang.

Cách dùng: python3 Report/apply-form-formatting.py <file.docx>
"""
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm

FONT = "Times New Roman"
BODY_PT = 13
CAPTION_PT = 12
TABLE_PT = 11
# Bề rộng vùng chữ của khổ A4 với lề của form mẫu, tính bằng twip
TABLE_WIDTH_TWIP = 9332
TITLE_PT = 13
TITLE_BIG_PT = 15
INDENT = Cm(1.0)

RE_TABLE_CAPTION = re.compile(r"^Bảng \d+\.\d+:")
RE_FIGURE_CAPTION = re.compile(r"^Hình \d+\.\d+:")
TITLE_BIG_MARKERS = ("NGHIÊN CỨU VÀ XÂY DỰNG", "SỬ DỤNG STABLE DIFFUSION")


def set_runs(paragraph, size_pt, bold=None, italic=None):
    """Đặt font, cỡ chữ và kiểu chữ cho mọi run của một đoạn."""
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(size_pt)
        # Ép cả font cho bảng mã Đông Á để Word không rơi về font mặc định
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic


def is_page_break(paragraph):
    """Đoạn chỉ chứa một ngắt trang, không có chữ."""
    return not paragraph.text.strip() and bool(
        paragraph._element.findall(".//" + qn("w:br"))
    )


def set_paragraph_mark_font(paragraph, size_pt):
    """Đặt font cho dấu kết đoạn, chính là cỡ chữ của số thứ tự hoặc dấu đầu dòng."""
    ppr = paragraph._element.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    sz = rpr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rpr.append(sz)
    sz.set(qn("w:val"), str(int(size_pt * 2)))


def is_list_item(paragraph):
    """Đoạn thuộc danh sách có đánh dấu đầu dòng hoặc đánh số."""
    return paragraph._element.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None


def normalize_undefined_styles(doc):
    """Gỡ các pStyle mà pandoc tham chiếu nhưng styles.xml không định nghĩa.

    File mẫu do người Việt soạn có styleId bị mã hoá thành 'a', 'a0'..., nên pandoc
    không ánh xạ được BodyText / FirstParagraph / Compact / ImageCaption. Để nguyên
    thì Word và LibreOffice dựng khác nhau, ví dụ dòng đầu trang bìa mất căn giữa.
    """
    defined = {style.style_id for style in doc.styles}
    fixed = 0
    for paragraph in doc.paragraphs:
        ppr = paragraph._element.find(qn("w:pPr"))
        if ppr is None:
            continue
        pstyle = ppr.find(qn("w:pStyle"))
        if pstyle is None:
            continue
        if pstyle.get(qn("w:val")) not in defined:
            ppr.remove(pstyle)
            fixed += 1
    return fixed


def has_image(paragraph):
    return bool(
        paragraph._element.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        )
    )


def build_toc_paragraph():
    """Dựng đoạn chứa trường TOC để Word tự sinh mục lục có số trang."""
    p = OxmlElement("w:p")

    def field_run(child):
        r = OxmlElement("w:r")
        r.append(child)
        p.append(r)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    field_run(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    field_run(instr)

    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    field_run(sep)

    placeholder = OxmlElement("w:t")
    placeholder.text = "Bấm phải chuột vào đây, chọn Update Field để hiện mục lục."
    field_run(placeholder)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run(end)
    return p


def replace_manual_toc(doc):
    """Xoá mục lục viết tay giữa 'MỤC LỤC' và 'DANH MỤC HÌNH VẼ', chèn trường TOC."""
    paragraphs = doc.paragraphs
    try:
        start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "MỤC LỤC")
        stop = next(
            i for i, p in enumerate(paragraphs) if p.text.strip() == "DANH MỤC HÌNH VẼ"
        )
    except StopIteration:
        return 0

    removed = 0
    for p in paragraphs[start + 1 : stop]:
        if is_page_break(p):
            continue
        p._element.getparent().remove(p._element)
        removed += 1

    paragraphs[start]._element.addnext(build_toc_paragraph())
    return removed


def format_title_page(doc):
    """Căn giữa và in đậm toàn bộ hai trang bìa, tới trước mục 'LỜI CẢM ƠN'."""
    count = 0
    for p in doc.paragraphs:
        if p.text.strip() == "LỜI CẢM ƠN":
            break
        if not p.text.strip():
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        size = (
            TITLE_BIG_PT
            if any(m in p.text for m in TITLE_BIG_MARKERS)
            else TITLE_PT
        )
        set_runs(p, size, bold=True)
        count += 1
    return count


def format_body(doc):
    """Áp định dạng thân bài, chú thích hình và chú thích bảng."""
    stats = {"body": 0, "caption": 0, "image": 0}
    in_title_page = True
    # Hai danh mục đầu sách cũng bắt đầu bằng "Hình x.y:" / "Bảng x.y:" nhưng là
    # dòng liệt kê, không phải chú thích, nên phải loại khỏi luật định dạng chú thích.
    in_index_list = False
    for p in doc.paragraphs:
        if p.text.strip() == "LỜI CẢM ƠN":
            in_title_page = False
        if in_title_page:
            continue
        if p.text.strip() == "DANH MỤC HÌNH VẼ":
            in_index_list = True
        elif p.text.strip() == "DANH MỤC TỪ VIẾT TẮT VÀ THUẬT NGỮ":
            in_index_list = False

        pf = p.paragraph_format
        if has_image(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
            stats["image"] += 1
            continue

        text = p.text.strip()
        if not text:
            continue

        if p.style.name.startswith("Heading"):
            pf.first_line_indent = None
            continue

        if RE_FIGURE_CAPTION.match(text) and not in_index_list:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = None
            pf.line_spacing = 1.0
            set_runs(p, CAPTION_PT, bold=True)
            stats["caption"] += 1
            continue

        if RE_TABLE_CAPTION.match(text) and not in_index_list:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = None
            pf.line_spacing = 1.0
            set_runs(p, CAPTION_PT, bold=True)
            stats["caption"] += 1
            continue

        # Đoạn văn thường: căn đều, giãn dòng 1,5, thụt đầu dòng 1cm.
        # Các dòng danh sách do pandoc sinh ra giữ nguyên lề trái, không thụt thêm.
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        if is_list_item(p):
            pf.first_line_indent = None
            set_paragraph_mark_font(p, BODY_PT)
        else:
            pf.first_line_indent = INDENT
        set_runs(p, BODY_PT)
        stats["body"] += 1
    return stats


def fix_table_look(doc):
    """Gắn lại kiểu bảng và kẻ khung.

    Pandoc ghi tblStyle "Table" nhưng styles.xml lấy từ form mẫu không có id này,
    nên bảng mất khung và mất bố cục. Chuyển sang id bảng có thật của form và kẻ
    khung trực tiếp cho chắc.
    """
    defined = {style.style_id for style in doc.styles}
    target = "a0" if "a0" in defined else None
    for table in doc.tables:
        tbl_pr = table._tbl.find(qn("w:tblPr"))
        if tbl_pr is None:
            continue
        tbl_style = tbl_pr.find(qn("w:tblStyle"))
        if tbl_style is not None:
            if target:
                tbl_style.set(qn("w:val"), target)
            else:
                tbl_pr.remove(tbl_style)
        if tbl_pr.find(qn("w:tblBorders")) is None:
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "000000")
                borders.append(el)
            tbl_pr.append(borders)
        # Cố định bề rộng bảng và bề rộng từng cột theo tỷ lệ pandoc đã tính.
        # Để kiểu "pct" thì LibreOffice dựng sai, chữ trong ô rơi ra ngoài bảng.
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(TABLE_WIDTH_TWIP))

        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")

        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            cols = grid.findall(qn("w:gridCol"))
            widths = [int(c.get(qn("w:w")) or 0) for c in cols]
            total = sum(widths)
            if cols:
                if total <= 0:
                    widths = [1] * len(cols)
                    total = len(cols)
                used = 0
                for i, col in enumerate(cols):
                    if i == len(cols) - 1:
                        w = TABLE_WIDTH_TWIP - used
                    else:
                        w = round(widths[i] * TABLE_WIDTH_TWIP / total)
                        used += w
                    col.set(qn("w:w"), str(w))
    return len(doc.tables)


def format_tables(doc):
    """Ô bảng dùng cỡ 11pt như form mẫu, không thụt đầu dòng, giãn dòng đơn."""
    cells = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = None
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    set_runs(p, TABLE_PT)
                cells += 1
    return cells


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "Report/Do-an-The-Data-Mint.docx"
    doc = docx.Document(path)

    normalized = normalize_undefined_styles(doc)
    removed = replace_manual_toc(doc)
    titles = format_title_page(doc)
    stats = format_body(doc)
    tables = fix_table_look(doc)
    cells = format_tables(doc)

    doc.save(path)
    print(f"Style chưa định nghĩa đã gỡ: {normalized} đoạn")
    print(f"Mục lục viết tay đã xoá: {removed} đoạn, đã chèn trường TOC")
    print(f"Trang bìa: {titles} đoạn căn giữa")
    print(f"Thân bài: {stats['body']} đoạn | chú thích: {stats['caption']} | hình: {stats['image']}")
    print(f"Bảng: {tables} | ô bảng: {cells}")
    print(f"Đã ghi đè: {path}")


if __name__ == "__main__":
    main()
